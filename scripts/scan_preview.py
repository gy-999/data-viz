
import sys, json, csv, gzip, time, re
from pathlib import Path
from collections import defaultdict
from datetime import datetime

def scan(source):
    """Unified scanner: auto-detect source type and scan."""
    t0 = time.time()
    
    # MySQL URL
    if source.startswith('mysql://'):
        return scan_mysql(source, t0)
    
    p = Path(source)
    if not p.exists():
        return {'status':'error','msg':'not found','source':source}
    
    # Directory
    if p.is_dir():
        return scan_directory(p, t0)
    
    name = p.name
    mb = p.stat().st_size / 1024 / 1024
    
    # Excel
    if name.endswith(('.xlsx','.xls')):
        import openpyxl
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        sheets = [{'name':s,'rows':max(ws.max_row or 0,0)} for s in wb.sheetnames for ws in [wb[s]]]
        wb.close()
        return {'type':'excel','name':name,'mb':round(mb,2),'sheets':sheets,'count':len(sheets),'status':'ok','scan_time':round(time.time()-t0,2)}
    
    # SQL script
    if name.endswith('.sql'):
        f = open(source, 'r', encoding='latin-1', errors='replace')
        c = f.read(1024*1024)  # Read first 1MB to detect
        f.seek(0)
        full = f.read() if len(c) < 100*1024 else None
        f.close()
        ins_count = full.upper().count('INSERT') if full else c.upper().count('INSERT')
        return {'type':'sql_script','name':name,'mb':round(mb,2),'total_inserts':ins_count,'is_large':mb>10,'status':'ok','scan_time':round(time.time()-t0,2)}
    
    # CSV/TSV/GZ
    is_gz = name.lower().endswith('.gz')
    fh = None
    for enc in ['utf-8-sig','utf-8','gbk','gb18030']:
        try:
            fh = (gzip.open if is_gz else open)(source, 'rt', encoding=enc)
            fh.readline()
            fh.seek(0)
            break
        except:
            if fh: fh.close()
    
    if not fh:
        return {'type':'text','name':name,'mb':round(mb,2),'status':'decode_fail','scan_time':round(time.time()-t0,2)}
    
    try:
        line = fh.readline()
        if not line:
            return {'type':'text','name':name,'mb':round(mb,2),'rows':0,'status':'empty','scan_time':round(time.time()-t0,2)}
        
        sc = [('|',line.count('|')),('\\t',line.count('\\t')),(',',line.count(',')),(';',line.count(';'))]
        sc.sort(key=lambda x:-x[1])
        sep = sc[0][0] if sc[0][1] > 0 else ','
        cols = [h.strip() for h in next(csv.reader([line.rstrip()], delimiter=sep))]
        
        total = 0
        samples = []
        for row in csv.reader(fh, delimiter=sep):
            total += 1
            if total <= 5000:
                samples.append(row)
            if total > 500000:
                break
        
        fh.close()
        
        return {
            'type':'text','name':name,'mb':round(mb,2),'rows':total,'cols':len(cols),'columns':cols,
            'sep':sep,'is_large':total>100000,'preview_is_sample':total>20,
            'preview':[{c:r[i].strip() if i<len(r) else '' for i,c in enumerate(cols)} for r in samples[:15]],
            'status':'ok','scan_time':round(time.time()-t0,2)
        }
    finally:
        try: fh.close()
        except: pass

def scan_mysql(source, t0):
    import pymysql
    # Parse mysql://user:pass@host:port/db
    rest = source[8:]  # strip mysql://
    u, rest = rest.split(':', 1)
    p, rest = rest.split('@', 1)
    if ':' in rest:
        h, rest = rest.split(':', 1)
        pt, db = rest.split('/', 1)
        pt = int(pt)
    else:
        h, db = rest.split('/', 1)
        pt = 3306
    pt = int(pt) if pt else 3306
    conn = pymysql.connect(host=h, user=u, password=p, database=db, connect_timeout=5)
    cur = conn.cursor()
    cur.execute('SELECT VERSION()')
    ver = cur.fetchone()[0]
    cur.execute('SHOW TABLES')
    allt = [r[0] for r in cur.fetchall()]
    tables = []
    trows = 0
    big = []
    for t in allt[:300]:
        cur.execute('SELECT COUNT(*) FROM ' + t)
        rc = cur.fetchone()[0]
        trows += rc
        cur.execute('SHOW COLUMNS FROM ' + t)
        cols = [{'n':c[0],'t':c[1]} for c in cur.fetchall()]
        tables.append({'name':t,'rows':rc,'col_count':len(cols)})
        if rc > 100000:
            big.append(t)
    cur.close()
    conn.close()
    return {
        'type':'mysql','database':db,'version':ver,'table_count':len(allt),
        'total_rows':trows,'tables':tables,'big_tables':big,
        'status':'ok','scan_time':round(time.time()-t0,2)
    }

def scan_directory(p, t0):
    entries = []
    by_type = defaultdict(list)
    for f in sorted(p.rglob('*')):
        if f.is_file():
            name = f.name.lower()
            if name.endswith('.csv'): t='csv'
            elif name.endswith('.tsv'): t='tsv'
            elif name.endswith(('.xlsx','.xls')): t='excel'
            elif name.endswith('.sql'): t='sql_script'
            elif name.endswith(('.csv.gz','.csv.gzip')): t='gz_csv'
            elif name.endswith('.json'): t='json'
            else: t='other'
            if t != 'other':
                entries.append({'type':t,'name':f.name,'mb':round(f.stat().st_size/1024/1024,2)})
                by_type[t].append(f.name)
    return {
        'type':'directory','name':p.name,'file_count':len(entries),
        'entries':entries,'by_type':dict(by_type),'status':'ok','scan_time':round(time.time()-t0,2)
    }

if __name__ == '__main__':
    import sys
    sources = [s for s in sys.argv[1:] if not s.startswith('--')]
    if not sources:
        print('Usage: python scan_preview.py <source1> [source2 ...]')
        sys.exit(1)
    if len(sources) == 1:
        r = scan(sources[0])
    else:
        r = {'multi':{s:scan(s) for s in sources},'count':len(sources)}
    print(json.dumps(r, ensure_ascii=False, default=str, indent=2))

