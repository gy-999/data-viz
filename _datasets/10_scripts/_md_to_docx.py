"""Convert the data-viz usage guide (Markdown) into a formatted Word document.

Handles: headings (1-3), paragraphs, bullet/numbered lists, pipe tables,
fenced code blocks, blockquotes, inline **bold** / `code`.
Sets CJK-safe fonts (eastAsia) so Chinese renders correctly.
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Twips

SRC = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path(
    r"D:\Claude cade\data\_datasets\data-viz-使用手册.md")
OUT = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else SRC.with_suffix(".docx")

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell_or_par, hexcolor):
    el = cell_or_par._element
    pr = el.get_or_add_tcPr() if el.tag.endswith("tc") else el.get_or_add_pPr()
    shd = pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): hexcolor})
    pr.append(shd)


def add_inline(par, text, base_size=10.5, base_bold=False):
    """Render **bold** and `code` spans."""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = par.add_run(tok[2:-2])
            set_run_font(run, BODY_FONT, base_size, True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            set_run_font(run, CODE_FONT, base_size - 0.5, base_bold, RGBColor(0xC0, 0x39, 0x2B))
        else:
            run = par.add_run(tok)
            set_run_font(run, BODY_FONT, base_size, base_bold)
    return par


def add_code_block(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.18)
        shade(p, "F4F5F7")
        run = p.add_run(line if line else " ")
        set_run_font(run, CODE_FONT, 9)
    # small gap after the block
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(4)
    set_run_font(tail.add_run(""), BODY_FONT, 4)


def _sum_cells(cells):
    total = 0
    for c in cells:
        for p in c.paragraphs:
            for r in p.runs:
                total += len(r.text)
    return max(total, 1)


def add_table(doc, rows, usable_twips=None):
    """Autofit tables can render wider than the text column and get clipped by
    viewers; pin the total width to the usable content width instead."""
    header, *body = rows
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = False

    # width budget proportional to header text length, so long labels get room
    weights = [max(len(h), 4) for h in header]
    wsum = sum(weights)
    budget = usable_twips or 9360
    widths = [max(int(budget * w / wsum), 700) for w in weights]
    scale = budget / sum(widths)
    widths = [int(w * scale) for w in widths]

    hdr = t.rows[0].cells
    for i, txt in enumerate(header):
        hdr[i].text = ""
        add_inline(hdr[i].paragraphs[0], txt, base_size=9.5, base_bold=True)
        shade(hdr[i], "EEF1F5")
    for row in body:
        cells = t.add_row().cells
        for i, txt in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], txt, base_size=9.5)

    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Twips(w)
    t.width = Twips(sum(widths))
    doc.add_paragraph()


def is_table_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def main():
    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x14, 0x56, 0xF0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
    usable = (doc.sections[0].page_width.twips - doc.sections[0].left_margin.twips
              - doc.sections[0].right_margin.twips)
    i = 0
    first_heading_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block)
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_table_sep(lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            add_table(doc, rows, usable_twips=usable)
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and not first_heading_done:
                h = doc.add_heading("", level=0)
                add_inline(h, text, base_size=22, base_bold=True)
                for r in h.runs:
                    set_run_font(r, BODY_FONT, 22, True, RGBColor(0x1F, 0x23, 0x29))
                first_heading_done = True
            else:
                h = doc.add_heading("", level=min(level, 3))
                add_inline(h, text, base_size={1: 17, 2: 14, 3: 12}.get(min(level, 3), 12), base_bold=True)
                for r in h.runs:
                    set_run_font(r, BODY_FONT, {1: 17, 2: 14, 3: 12}.get(min(level, 3), 12), True,
                                 RGBColor(0x14, 0x56, 0xF0))
            i += 1
            continue

        # horizontal rule
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("─" * 46)
            set_run_font(run, BODY_FONT, 8, False, RGBColor(0xAE, 0xB4, 0xBD))
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            shade(p, "F7F8FA")
            add_inline(p, stripped.lstrip("> ").strip(), base_size=10)
            i += 1
            continue

        # bullets / numbered
        m = re.match(r"^([-*])\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"written: {OUT}  ({OUT.stat().st_size/1024:.0f} KB)")
    print(f"paragraphs: {len(doc.paragraphs)}  tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
